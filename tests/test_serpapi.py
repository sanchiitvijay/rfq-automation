"""SerpApi first, ferguson.com page as fallback, the rate limit, supplier info and source links."""
from __future__ import annotations

import asyncio
import json
import time

import httpx
import openpyxl
from helpers import FIXTURES  # noqa: F401  (fixtures dir exists)

from rfq.config import RunOptions
from rfq.eta import Progress
from rfq.excel.writer import render
from rfq.extract_ids import extract
from rfq.lookup import resolve as resolve_mod
from rfq.lookup.cache import MemoryCache
from rfq.lookup.live import LiveResult
from rfq.lookup.serpapi import EXHAUSTED_FLAG, parse_response
from rfq.models import Cell

PAGE = "https://www.ferguson.com/product/presco-6-in-x-1000-ft-blue-detector-tape-psd6105b52/5106912.html"


def serp_ok(pid: str, url: str, title: str = "Presco 6 in x 1000 ft Blue Detector Tape | Ferguson") -> dict:
    return {"search_metadata": {"status": "Success"}, "organic_results": [
        {"position": 1, "title": "Detector tape category | Ferguson",
         "link": "https://www.ferguson.com/category/marking-tape/"},  # not the product: must be ignored
        {"position": 2, "title": title, "link": url + "?srsltid=abc",
         "rich_snippet": {"bottom": {"extensions": ["$139.02", "In stock"],
                                     "detected_extensions": {"price": 139.02}}}},
    ]}


def test_parse_found_with_price_and_stock():
    r = parse_response(serp_ok("5106912", PAGE), "5106912", read_on="2026-10-16")
    assert r.status == "found" and r.url == PAGE and r.title == "Presco 6 in x 1000 ft Blue Detector Tape"
    assert r.supplier["Price (USD)"].value == "139.02" and r.supplier["Availability"].value == "In stock"
    assert all(c.url == PAGE and c.source_type == "serpapi" for c in r.supplier.values())


def test_parse_ignores_other_products_and_sites():
    data = {"organic_results": [
        {"title": "x", "link": "https://www.ferguson.com/product/other/5106935.html"},
        {"title": "y", "link": "https://www.homedepot.com/p/5106912"},
    ]}
    assert parse_response(data, "5106912").status == "no_result"


def test_parse_truncated_title_not_used_as_description():
    r = parse_response(serp_ok("5106912", PAGE, title="Presco 6 in x 1000 ft Blue Detector Tape CAUTION ..."), "5106912")
    assert r.title_truncated


def test_parse_errors():
    assert parse_response({"error": "Google hasn't returned any results for this query."}, "1").status == "no_result"
    assert parse_response({"error": "Your account has run out of searches."}, "1").status == "exhausted"
    assert parse_response({"error": "Invalid API key. Your API key should be here: …"}, "1").status == "exhausted"
    assert parse_response({"error": "Something odd"}, "1").status == "error"


class FakePages:
    results: dict[str, LiveResult] = {}
    requested: list[str] = []
    waited_s = 0.0

    def __init__(self, *a, **k): ...
    async def __aenter__(self): return self
    async def __aexit__(self, *exc): return None

    async def lookup(self, pid):
        FakePages.requested.append(pid)
        return self.results.get(pid, LiveResult("not_found", "not on ferguson.com"))


def run_resolve(monkeypatch, body: str, serp_handler, pages=None, key="test-key", policy="missing", cache=None):
    FakePages.results, FakePages.requested = pages or {}, []
    monkeypatch.setattr(resolve_mod, "LiveFerguson", FakePages)
    calls: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        assert "api_key=test-key" in str(request.url)
        q = request.url.params["q"]
        calls.append(q)
        status, data = serp_handler(q.split('"')[1])
        return httpx.Response(status, json=data)

    real = httpx.AsyncClient
    monkeypatch.setattr(resolve_mod.httpx, "AsyncClient", lambda *a, **k: real(transport=httpx.MockTransport(handler)))
    opts = RunOptions(offline=False, use_llm=False, use_pdfs=False, serpapi_key=key, scrape_after_serpapi=policy)
    cache = cache or MemoryCache()
    sheet = asyncio.run(resolve_mod.resolve(extract(body).items, opts, cache, Progress()))
    return sheet, calls, cache


def test_serpapi_hit_in_catalog_skips_ferguson_page(monkeypatch):
    sheet, calls, _ = run_resolve(monkeypatch, "5106912", lambda pid: (200, serp_ok(pid, PAGE)))
    row = sheet.rows[0]
    assert calls and FakePages.requested == []  # no ferguson.com page load at all
    assert row.lookup_status == "Found" and row.lookup_source == "serpapi+catalog"
    assert row.value("Weight") == "14.25 lb"  # specs from the catalog
    assert row.supplier["Price (USD)"].source_type == "serpapi"  # fresher price from Google beats the catalog
    assert any(t.startswith("SerpApi (Google): found") for t in row.tried)


def test_serpapi_hit_not_in_catalog_reads_the_page(monkeypatch):
    url = "https://www.ferguson.com/product/x/225123.html"
    page = LiveResult("found", cells={"Width": Cell("6 in", "ferguson_page", "high", "page", url)},
                      supplier={"Brand": Cell("Pro-Pak", "ferguson_page", "high", "page", url)})
    sheet, _, _ = run_resolve(monkeypatch, "225123", lambda pid: (200, serp_ok(pid, url)), {"225123": page})
    row = sheet.rows[0]
    assert FakePages.requested == ["225123"] and row.lookup_source == "live" and row.value("Width") == "6 in"
    assert row.supplier["Brand"].value == "Pro-Pak" and row.supplier["Price (USD)"].source_type == "serpapi"


def test_serpapi_hit_page_fails_not_in_catalog_keeps_identity_only(monkeypatch):
    url = "https://www.ferguson.com/product/x/225123.html"
    sheet, _, _ = run_resolve(monkeypatch, "225123", lambda pid: (200, serp_ok(pid, url, "6 in. Water Tape | Ferguson")),
                              {"225123": LiveResult("blocked", "HTTP 403")})
    row = sheet.rows[0]
    assert row.lookup_status == "Found" and row.lookup_source == "serpapi"
    assert row.value("Description") == "6 in. Water Tape" and row.cells["Description"].url == url
    assert not row.value("Width")  # never guessed
    assert any("specs left blank" in n.note for n in sheet.notes)


def test_serpapi_quota_exhausted_falls_back_to_scraping(monkeypatch):
    sheet, calls, cache = run_resolve(monkeypatch, "5106912\n5106935",
                                      lambda pid: (429, {"error": "Your account has run out of searches."}))
    assert set(FakePages.requested) == {"5106912", "5106935"}
    assert asyncio.run(cache.has_flag(EXHAUSTED_FLAG))  # switched off for the next jobs too
    assert any("SerpApi stopped" in w for w in sheet.warnings)
    assert all(r.lookup_status == "Found" for r in sheet.rows)  # catalog after the failed page reads


def test_serpapi_not_configured_scrapes(monkeypatch):
    sheet, calls, _ = run_resolve(monkeypatch, "5106912", lambda pid: (200, {}), key="")
    assert calls == [] and FakePages.requested == ["5106912"]
    assert any("SERPAPI_API_KEY is not set" in w for w in sheet.warnings)


def test_serpapi_no_result_still_checks_the_page(monkeypatch):
    sheet, _, _ = run_resolve(monkeypatch, "9677533",
                              lambda pid: (200, {"error": "Google hasn't returned any results for this query."}))
    row = sheet.rows[0]
    assert FakePages.requested == ["9677533"] and row.lookup_status == "Not found"
    assert "Google (SerpApi) has no ferguson.com page" in row.not_found_reason


def test_serpapi_results_are_cached(monkeypatch):
    cache = MemoryCache()
    _, first, _ = run_resolve(monkeypatch, "5106912", lambda pid: (200, serp_ok(pid, PAGE)), cache=cache)
    _, second, _ = run_resolve(monkeypatch, "5106912", lambda pid: (200, serp_ok(pid, PAGE)), cache=cache)
    assert len(first) == 1 and second == []


def test_policy_always_reads_page_even_after_hit(monkeypatch):
    run_resolve(monkeypatch, "5106912", lambda pid: (200, serp_ok(pid, PAGE)), policy="always")
    assert FakePages.requested == ["5106912"]


def test_rate_limit_caps_calls_per_window():
    async def go():
        cache = MemoryCache()
        t0 = time.monotonic()
        stamps = []

        async def one():
            await cache.rate("ferguson", 12, 1.0)  # 12 per second here, 12 per minute in production
            stamps.append(time.monotonic() - t0)

        await asyncio.gather(*(one() for _ in range(30)))
        return sorted(stamps)

    stamps = asyncio.run(go())
    assert all(sum(1 for s in stamps if a <= s < a + 0.99) <= 12 for a in stamps)
    assert stamps[-1] >= 2.0  # 30 calls at 12 per window need at least 3 windows


def test_supplier_sheet_and_source_links(tmp_path, monkeypatch):
    sheet, _, _ = run_resolve(monkeypatch, "5106912\nMTX-4409-XQ", lambda pid: (200, serp_ok(pid, PAGE)))
    path = render(sheet, tmp_path / "RFQ.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames.index("Supplier") == wb.sheetnames.index("RFQ") + 1
    sup = list(wb["Supplier"].iter_rows(values_only=True))
    head = sup[0]
    first = dict(zip(head, sup[1]))
    assert first["Supplier"] == "Ferguson" and first["Brand"] == "Presco" and first["Price (USD)"] == "139.02"
    assert first["Found via"] == "SerpApi + backup catalog"
    assert dict(zip(head, sup[2]))["Lookup Status"] == "Not found" and dict(zip(head, sup[2]))["Price (USD)"] is None
    links = [c for r in wb["Supplier"].iter_rows() for c in r if c.hyperlink]
    assert links and all(c.hyperlink.target.startswith("https://") for c in links)
    src = wb["Sources"]
    assert [c.value for c in src[1]][:7] == ["product_id", "column", "value", "source_type", "confidence", "source", "link"]
    rows = list(src.iter_rows(min_row=2))
    weight = next(r for r in rows if r[1].value == "Weight")
    assert weight[6].hyperlink and "ferguson.com/product" in weight[6].value
    price = next(r for r in rows if r[1].value == "Supplier: Price (USD)")
    assert price[3].value == "serpapi" and price[6].value == PAGE
    rule = next(r for r in rows if r[3].value == "validated_rule")
    assert rule[6].value is None  # a rule has no single page: no made-up link


def test_formula_like_link_never_becomes_a_link(tmp_path):
    from rfq.models import Row, Sheet

    row = Row(product_id="5106912", kind="ferguson", lookup_status="Found")
    row.cells["Description"] = Cell("x", "manual_edit", "high", "me", 'javascript:alert(1)')
    path = render(Sheet(rows=[row]), tmp_path / "x.xlsx")
    cell = openpyxl.load_workbook(path)["Sources"]["G2"]
    assert cell.value is None and cell.hyperlink is None


def test_reader_round_trips_supplier_and_links(tmp_path, monkeypatch):
    from rfq.excel.reader import read

    sheet, _, _ = run_resolve(monkeypatch, "5106912", lambda pid: (200, serp_ok(pid, PAGE)))
    back = read(render(sheet, tmp_path / "RFQ.xlsx"))
    row = back.rows[0]
    assert row.supplier["Price (USD)"].value == "139.02" and row.supplier["Price (USD)"].url == PAGE
    assert row.cells["Weight"].url.startswith("https://www.ferguson.com/product/")
    json.dumps(back.to_dict())  # still serialisable for the backend
