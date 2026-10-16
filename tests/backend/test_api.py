from __future__ import annotations

import asyncio
import base64

import pytest
from backend_helpers import drain

EMAILS = __import__("pathlib").Path(__file__).resolve().parents[2] / "emails"
pytestmark = pytest.mark.asyncio


def raw(name: str) -> str:
    return base64.b64encode(next(EMAILS.glob(f"{name}*.eml")).read_bytes()).decode()


async def make_job(client, queue, ctx, name="email-3", **extra):
    r = await client.post("/api/jobs", json={"source": "extension", "raw_eml_b64": raw(name), **extra})
    assert r.status_code == 201, r.text
    await drain(queue, ctx)
    return r.json()["job_id"]


async def test_auth_required(env):
    client, *_ = env
    assert (await client.get("/api/jobs", headers={"X-API-Key": "nope"})).status_code == 401


async def test_create_run_and_detail(env):
    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx)
    d = (await client.get(f"/api/jobs/{jid}")).json()
    assert d["status"] == "needs_review" and d["version"] == 1
    assert d["requested"] == 7 and d["not_found"] == 3
    assert [i["product_id"] for i in d["items"]][-1] == "MTX-4409-XQ"
    assert any(e["phase"] == "Done" for e in d["events"])
    x = await client.get(f"/api/jobs/{jid}/rfq.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"


async def test_gmail_id_is_idempotent(env):
    client, queue, ctx, _ = env
    body = {"source": "extension", "raw_eml_b64": raw("email-1"), "gmail_msg_id": "msg-f:1"}
    a = (await client.post("/api/jobs", json=body)).json()
    b = (await client.post("/api/jobs", json=body)).json()
    assert a["job_id"] == b["job_id"] and b["created"] is False and len(queue.jobs) == 1
    c = (await client.post("/api/jobs", json={**body, "force": True})).json()
    assert c["created"] and c["job_id"] != a["job_id"]


async def test_bad_inputs(env):
    client, *_ = env
    assert (await client.post("/api/jobs", json={"source": "paste", "body_text": "  "})).status_code == 422
    assert (await client.post("/api/jobs", json={"source": "extension", "raw_eml_b64": base64.b64encode(b"no headers here").decode()})).status_code == 422
    r = await client.post("/api/jobs/upload", files={"file": ("x.txt", b"hello", "text/plain")})
    assert r.status_code == 422


async def test_paste_job(env):
    client, queue, ctx, _ = env
    r = await client.post("/api/jobs", json={"source": "paste", "subject": "RFQ", "sender": "a@b.com", "body_text": "5106912\n5106935"})
    await drain(queue, ctx)
    d = (await client.get(f"/api/jobs/{r.json()['job_id']}")).json()
    assert d["found"] == 2 and d["email"]["has_raw"] is False


async def test_edit_validation_and_locking(env):
    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx)
    s = (await client.get(f"/api/jobs/{jid}/sheet")).json()
    pid = next(r["product_id"] for r in s["rows"] if r["lookup_status"] == "Found" and "Weight" not in r["cells"])
    ok = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 1, "edits": [
        {"product_id": pid, "column": "Weight", "value": "12.5 lb", "note": "quote"}]})
    assert ok.status_code == 200 and ok.json()["version"] == 2
    cell = next(r for r in ok.json()["rows"] if r["product_id"] == pid)["cells"]["Weight"]
    assert cell["source_type"] == "manual_edit" and "Tester" in cell["source"]
    stale = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 1, "edits": [
        {"product_id": pid, "column": "Weight", "value": "13 lb"}]})
    assert stale.status_code == 409 and stale.json()["current_version"] == 2
    a, b = await asyncio.gather(*[client.patch(f"/api/jobs/{jid}/sheet", json={"version": 2, "edits": [
        {"product_id": pid, "column": "Weight", "value": v}]}) for v in ("13 lb", "14 lb")])
    assert sorted([a.status_code, b.status_code]) == [200, 409]
    for bad in ("heavy", "=1+1", "12"):
        r = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 3, "edits": [
            {"product_id": pid, "column": "Weight", "value": bad}]})
        assert r.status_code == 422, bad
    r = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 3, "edits": [
        {"product_id": pid, "column": "product_id", "value": "1"}]})
    assert r.status_code == 422


async def test_fill_not_found_and_remove_row(env):
    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx)
    r = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 1, "edits": [
        {"product_id": "MTX-4409-XQ", "column": "Description", "value": "Matched by phone", "note": "customer"}],
        "remove_rows": [{"product_id": "31-024", "reason": "dropped"}]})
    rows = {x["product_id"]: x for x in r.json()["rows"]}
    assert rows["MTX-4409-XQ"]["lookup_status"] == "Manual" and "31-024" not in rows
    d = (await client.get(f"/api/jobs/{jid}")).json()
    assert any("Removed from the RFQ by Tester: dropped" in n["note"] for n in d["notes"])


async def test_send_flow(env):
    client, queue, ctx, sent = env
    jid = await make_job(client, queue, ctx)
    base = {"to": "priya.nair@sourcing.example.com", "subject": "Re: RFQ", "body": "attached",
            "idempotency_key": "key-00000001"}
    assert (await client.post(f"/api/jobs/{jid}/send", json={**base, "to": "nope"})).status_code == 422
    incomplete = await client.post(f"/api/jobs/{jid}/send", json=base)
    assert incomplete.status_code == 409
    a, b = await asyncio.gather(*[client.post(f"/api/jobs/{jid}/send", json={**base, "confirm_incomplete": True})
                                  for _ in range(2)])
    assert {a.status_code, b.status_code} == {202} and a.json()["id"] == b.json()["id"]
    await drain(queue, ctx)
    assert len(sent) == 1
    msg = sent[0]
    assert msg["To"] == "priya.nair@sourcing.example.com" and msg["In-Reply-To"] == "<rfq-sample-3@sourcing.example.com>"
    assert [p.get_filename() for p in msg.iter_attachments()] == ["RFQ.xlsx"]
    d = (await client.get(f"/api/jobs/{jid}")).json()
    assert d["status"] == "sent" and d["sends"][0]["status"] == "sent"
    again = await client.post(f"/api/jobs/{jid}/send", json={**base, "idempotency_key": "key-00000002", "confirm_incomplete": True})
    assert again.status_code == 409
    resend = await client.post(f"/api/jobs/{jid}/send", json={**base, "idempotency_key": "key-00000003",
                                                              "confirm_incomplete": True, "resend": True})
    assert resend.status_code == 202
    await drain(queue, ctx)
    assert len(sent) == 2


async def test_send_permanent_failure_returns_to_review(env, monkeypatch):
    client, queue, ctx, _ = env
    from backend.app import worker as worker_mod
    from backend.app.services.mail import PermanentSendError

    def fail(settings, msg):
        raise PermanentSendError("Gmail login failed")

    monkeypatch.setattr(worker_mod, "smtp_send", fail)
    jid = await make_job(client, queue, ctx)
    await client.post(f"/api/jobs/{jid}/send", json={"to": "a@example.com", "subject": "s", "body": "b",
                                                     "idempotency_key": "key-fail-0001", "confirm_incomplete": True})
    await drain(queue, ctx)
    d = (await client.get(f"/api/jobs/{jid}")).json()
    assert d["status"] == "needs_review" and "Gmail login failed" in d["error"] and d["sends"][0]["status"] == "failed"


async def test_recipient_domain_allowlist(env, monkeypatch):
    client, queue, ctx, _ = env
    from backend.app.settings import get_settings

    monkeypatch.setenv("ALLOWED_RECIPIENT_DOMAINS", "sourcing.example.com")
    get_settings.cache_clear()
    jid = await make_job(client, queue, ctx)
    r = await client.post(f"/api/jobs/{jid}/send", json={"to": "x@evil.example.org", "subject": "s", "body": "b",
                                                         "idempotency_key": "key-dom-00001", "confirm_incomplete": True})
    assert r.status_code == 422 and "not an allowed" in r.json()["detail"]


async def test_worker_failure_retries_then_fails(env, monkeypatch):
    client, queue, ctx, _ = env
    from arq import Retry

    from backend.app import worker as worker_mod

    async def boom(*a, **k):
        raise RuntimeError("disk full")

    monkeypatch.setattr(worker_mod, "run", boom)
    r = await client.post("/api/jobs", json={"source": "paste", "body_text": "5106912"})
    jid = r.json()["job_id"]
    name, args = queue.jobs.pop(0)
    with pytest.raises(Retry):
        await worker_mod.run_rfq_job({**ctx, "job_try": 1}, *args)
    assert (await client.get(f"/api/jobs/{jid}")).json()["status"] == "queued"
    await worker_mod.run_rfq_job({**ctx, "job_try": 2}, *args)
    d = (await client.get(f"/api/jobs/{jid}")).json()
    assert d["status"] == "failed" and "disk full" in d["error"]
    assert (await client.post(f"/api/jobs/{jid}/retry")).status_code == 200


async def test_delete_job(env):
    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx, name="email-1")
    assert (await client.delete(f"/api/jobs/{jid}")).status_code == 204
    assert (await client.get(f"/api/jobs/{jid}")).status_code == 404


async def test_job_saved_without_queue_is_recovered_later(env):
    import uuid
    from datetime import datetime, timedelta, timezone

    from sqlalchemy import update

    from backend.app import db
    from backend.app.models import Job
    from backend.app import worker as worker_mod
    from backend.app.main import app

    client, queue, ctx, _ = env
    app.state.queue = None  # Redis down when the email arrives
    r = await client.post("/api/jobs", json={"source": "paste", "body_text": "5106912"})
    assert r.status_code == 201 and r.json()["queued"] is False and "starts automatically" in r.json()["message"]
    jid = r.json()["job_id"]
    async with db.sessions()() as s:  # pretend it has been waiting for 2 minutes
        await s.execute(update(Job).where(Job.id == uuid.UUID(jid))
                        .values(updated_at=datetime.now(timezone.utc) - timedelta(minutes=2)))
        await s.commit()
    assert await worker_mod.recover_stuck_jobs({**ctx, "redis": queue}) == "requeued 1"
    await drain(queue, ctx)
    assert (await client.get(f"/api/jobs/{jid}")).json()["status"] == "needs_review"


async def test_job_queued_twice_runs_once(env):
    from backend.app import worker as worker_mod

    client, queue, ctx, _ = env
    r = await client.post("/api/jobs", json={"source": "paste", "body_text": "5106912"})
    name, args = queue.jobs.pop(0)
    first, second = await asyncio.gather(worker_mod.run_rfq_job(ctx, *args), worker_mod.run_rfq_job(ctx, *args))
    assert sorted([first, second]) == ["done", "skipped (already taken, finished or deleted)"]
    d = (await client.get(f"/api/jobs/{r.json()['job_id']}")).json()
    assert d["status"] == "needs_review" and d["version"] == 1


async def test_import_preview_endpoint_then_apply(env):
    import io

    import openpyxl

    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx, name="email-1")
    xlsx = (await client.get(f"/api/jobs/{jid}/rfq.xlsx")).content
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames[0] == "RFQ"
    ws = wb["RFQ"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "5106880":
            row[header.index("ANSI")].value = "ANSI Z535.1"
            row[header.index("Weight")].value = "heavy"
    buf = io.BytesIO()
    wb.save(buf)
    r = await client.post(f"/api/jobs/{jid}/sheet/import", files={"file": ("RFQ.xlsx", buf.getvalue(), "application/octet-stream")})
    p = r.json()
    assert r.status_code == 200 and p["ok"] and p["file_version"] == 1 and p["counts"] == {"ok": 1, "invalid": 1, "conflict": 0}
    assert (await client.get(f"/api/jobs/{jid}")).json()["version"] == 1  # preview wrote nothing
    ok = [c for c in p["changes"] if c["status"] == "ok"]
    applied = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 1, "edits": [
        {"product_id": c["product_id"], "column": c["column"], "value": c["uploaded"], "note": "Uploaded Excel 'RFQ.xlsx'"}
        for c in ok]})
    assert applied.status_code == 200
    cell = next(x for x in applied.json()["rows"] if x["product_id"] == "5106880")["cells"]["ANSI"]
    assert cell["value"] == "ANSI Z535.1" and "Uploaded Excel" in cell["source"]


async def test_recheck_never_stores_a_report_for_an_old_version(env, monkeypatch):
    """A slow AI check that overlaps an edit must end with a report for the newest sheet."""
    from backend.app.services import sheet as sheet_svc
    from rfq.quality import agent

    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx, name="email-1")
    s = (await client.get(f"/api/jobs/{jid}/sheet")).json()
    pid = "5106880"
    assert any(f["product_id"] == pid and f["column"] == "ANSI" and f["type"] == "missing" for f in s["flags"])

    real = agent.run_quality
    started = asyncio.Event()
    calls = {"n": 0}

    async def slow(sheet, opts, cache=None, llm_client=None, llm=None, progress=None, llm_cache_only=False):
        if llm:
            calls["n"] += 1
            if calls["n"] == 1:
                started.set()
                await asyncio.sleep(0.6)  # the AI is "thinking" while someone saves an edit
        return await real(sheet, opts, cache, llm_client, llm=False, progress=progress)

    monkeypatch.setattr(sheet_svc, "run_quality", slow)

    async def edit_meanwhile():
        await started.wait()
        r = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 1, "edits": [
            {"product_id": pid, "column": "ANSI", "value": "ANSI Z535.1", "note": "x"}]})
        assert r.status_code == 200

    recheck, _ = await asyncio.gather(client.post(f"/api/jobs/{jid}/recheck"), edit_meanwhile())
    assert recheck.status_code == 200 and recheck.json()["version"] == 2 and calls["n"] == 2  # re-ran on v2
    final = (await client.get(f"/api/jobs/{jid}/sheet")).json()
    assert not any(f["product_id"] == pid and f["column"] == "ANSI" for f in final["flags"])  # no stale grey flag


async def test_review_is_recorded_and_reopened_by_edits(env):
    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx, name="email-1")
    assert (await client.post(f"/api/jobs/{jid}/review?done=true&version=0")).status_code == 409  # stale view
    assert (await client.post(f"/api/jobs/{jid}/review?done=true&version=1")).json()["status"] == "ready"
    await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 1, "edits": [
        {"product_id": "5106880", "column": "ANSI", "value": "ANSI Z535.1"}]})
    d = (await client.get(f"/api/jobs/{jid}")).json()
    assert d["status"] == "needs_review"
    phases = [e["phase"] for e in d["events"]]
    assert "Reviewed" in phases and phases[-1] == "Review reopened"
    assert any("by Tester" in e["message"] for e in d["events"] if e["phase"] == "Reviewed")


async def test_mail_test_endpoint_reports_login_problem(env, monkeypatch):
    from backend.app.services import mail

    client, *_ = env

    def reject(settings):
        raise mail.PermanentSendError("mail login rejected (535): Username and Password not accepted")

    monkeypatch.setattr(mail, "check_login", reject)
    r = (await client.post("/api/mail/test")).json()
    assert r["ok"] is False and "535" in r["detail"]


async def test_search_filters_pagination_and_suggestions(env):
    from datetime import datetime, timedelta, timezone

    client, queue, ctx, _ = env
    ids = {}
    for name in ("email-1", "email-2", "email-3", "email-4"):
        ids[name] = await make_job(client, queue, ctx, name=name)
    r = await client.post("/api/jobs", json={"source": "paste", "subject": "100% urgent_quote", "sender": "Ann <ann@x.com>",
                                             "body_text": "5106912"})
    await drain(queue, ctx)

    def items(resp):
        assert resp.status_code == 200, resp.text
        return resp.json()

    all_ = items(await client.get("/api/jobs"))
    assert all_["total"] == 5 and all_["page"] == 1 and all_["pages"] == 1
    assert [j["subject"] for j in all_["items"]][0] == "100% urgent_quote"  # newest first
    # free text matches subject, sender, body and product IDs
    assert items(await client.get("/api/jobs", params={"q": "MTX-4409"}))["total"] == 1
    assert items(await client.get("/api/jobs", params={"q": "barrier tapes"}))["total"] == 1
    assert items(await client.get("/api/jobs", params={"product": "5106912"}))["total"] == 2
    assert items(await client.get("/api/jobs", params={"sender": "ann@x.com"}))["total"] == 1
    assert items(await client.get("/api/jobs", params={"subject": "fwd"}))["total"] == 1
    # % and _ are literal characters, not wildcards
    assert items(await client.get("/api/jobs", params={"q": "100%"}))["total"] == 1
    assert items(await client.get("/api/jobs", params={"q": "%"}))["total"] == 1
    assert items(await client.get("/api/jobs", params={"q": "t_pes"}))["total"] == 0
    # status / source, repeated or comma separated
    assert items(await client.get("/api/jobs", params={"status": "needs_review"}))["total"] == 5
    assert items(await client.get("/api/jobs", params=[("source", "paste"), ("source", "extension")]))["total"] == 5
    assert items(await client.get("/api/jobs", params={"source": "paste"}))["total"] == 1
    # dates (timezone-aware), and a bad range
    now = datetime.now(timezone.utc)
    assert items(await client.get("/api/jobs", params={"created_from": (now - timedelta(hours=1)).isoformat()}))["total"] == 5
    assert items(await client.get("/api/jobs", params={"created_to": (now - timedelta(days=1)).isoformat()}))["total"] == 0
    bad = await client.get("/api/jobs", params={"created_from": now.isoformat(), "created_to": (now - timedelta(days=1)).isoformat()})
    assert bad.status_code == 422 and "after" in bad.text
    assert (await client.get("/api/jobs", params={"created_from": "yesterday"})).status_code == 422
    # pagination and bounds
    p2 = items(await client.get("/api/jobs", params={"page_size": 2, "page": 2, "sort": "oldest"}))
    assert p2["pages"] == 3 and len(p2["items"]) == 2 and p2["page"] == 2
    past = items(await client.get("/api/jobs", params={"page_size": 2, "page": 99}))
    assert past["page"] == 3 and len(past["items"]) == 1  # past the end shows the last page
    assert (await client.get("/api/jobs", params={"page_size": 1000})).status_code == 422
    assert (await client.get("/api/jobs", params={"sort": "drop table"})).status_code == 422
    assert (await client.get("/api/jobs", params={"status": "bogus"})).status_code == 422
    assert (await client.get("/api/jobs", params={"q": "x" * 201})).status_code == 422
    # facets come from the server (labels + counts), suggestions from the data
    f = items(await client.get("/api/jobs/facets"))
    assert {s["value"]: s["count"] for s in f["statuses"]}["needs_review"] == 5
    assert {s["value"] for s in f["sources"]} >= {"paste", "extension"} and f["page_sizes"] and f["sorts"]
    sug = (await client.get("/api/jobs/suggest", params={"q": "5106"})).json()
    assert sug and all(s["field"] == "product" for s in sug) and sug[0]["value"].startswith("5106")
    assert any(s["field"] == "sender" for s in (await client.get("/api/jobs/suggest", params={"q": "priya"})).json())
    assert (await client.get("/api/jobs/suggest", params={"q": "a"})).json() == []  # too short


async def test_summary_download_and_image_follows_picture(env):
    import io

    import openpyxl

    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx, name="email-1")
    r = await client.get(f"/api/jobs/{jid}/summary.xlsx")
    assert r.status_code == 200 and "RFQ_summary.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Notes"]
    assert (await client.get("/api/jobs/00000000-0000-0000-0000-000000000000/summary.xlsx")).status_code == 404
    bad = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 1, "edits": [
        {"product_id": "5106912", "column": "Image", "value": "https://i.ferg-img.com/i/fergusonprod/x_primary"}]})
    assert bad.status_code == 422 and "follows Picture" in bad.text
    ok = await client.patch(f"/api/jobs/{jid}/sheet", json={"version": 1, "edits": [
        {"product_id": "5106912", "column": "Picture", "value": "https://i.ferg-img.com/i/fergusonprod/x_primary"}]})
    row = next(x for x in ok.json()["rows"] if x["product_id"] == "5106912")
    assert row["cells"]["Image"]["value"] == row["cells"]["Picture"]["value"] == "https://i.ferg-img.com/i/fergusonprod/x_primary"


async def test_timestamps_are_utc_aware(env):
    client, queue, ctx, _ = env
    jid = await make_job(client, queue, ctx, name="email-1")
    d = (await client.get(f"/api/jobs/{jid}")).json()
    assert d["created_at"].endswith(("+00:00", "Z")) and d["events"][0]["ts"].endswith(("+00:00", "Z"))
