"""Upload of an edited RFQ workbook: validated, three-way compared, never applied blindly."""
from __future__ import annotations

import io

import openpyxl
from helpers import eml, run_sync

from rfq.excel.importer import preview
from rfq.excel.writer import render


def _download(tmp_path, version="1", job_id="job-1"):
    result = run_sync(tmp_path / "run", eml("email-1"))
    path = render(result.sheet, tmp_path / "RFQ_v1.xlsx", result.report, meta={"job_id": job_id, "version": version},
                  embed_images=False)
    return result.sheet, path


def _edit(path, edits: dict[tuple[str, str], object], drop_meta=False, rename=None) -> bytes:
    wb = openpyxl.load_workbook(path)
    ws = wb["RFQ"]
    header = [c.value for c in ws[1]]
    for (pid, column), value in edits.items():
        for row in ws.iter_rows(min_row=2):
            if str(row[0].value) == pid:
                row[header.index(column)].value = value
    if drop_meta:
        del wb["_rfq_meta"]
    if rename:
        ws.cell(row=1, column=header.index(rename[0]) + 1, value=rename[1])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_opens_on_rfq_sheet(tmp_path):
    _, path = _download(tmp_path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[0] == "RFQ" and wb.active.title == "RFQ"
    assert wb["_rfq_meta"].sheet_state == "hidden"


def test_unchanged_file_has_no_changes(tmp_path):
    sheet, path = _download(tmp_path)
    p = preview(path.read_bytes(), "RFQ.xlsx", "job-1", sheet, 1)
    assert p.ok and p.changes == [] and not p.errors


def test_valid_edits_and_invalid_ones(tmp_path):
    sheet, path = _download(tmp_path)
    data = _edit(path, {("5106880", "ANSI"): "ANSI Z535.1", ("5106880", "Weight"): "heavy",
                        ("5106912", "No of Tools"): 2, ("5106935", "Thickness"): "=1+1"})
    p = preview(data, "RFQ.xlsx", "job-1", sheet, 1)
    by = {(c.product_id, c.column): c for c in p.changes}
    assert by[("5106880", "ANSI")].status == "ok" and by[("5106880", "ANSI")].uploaded == "ANSI Z535.1"
    assert by[("5106880", "Weight")].status == "invalid" and "unit" in by[("5106880", "Weight")].message
    assert by[("5106912", "No of Tools")].status == "ok" and by[("5106912", "No of Tools")].uploaded == "2"
    assert by[("5106935", "Thickness")].status == "invalid" and "formulas" in by[("5106935", "Thickness")].message
    assert p.to_dict()["counts"] == {"ok": 2, "invalid": 2, "conflict": 0}


def test_image_formula_from_excel_365_is_not_a_change(tmp_path):
    sheet, path = _download(tmp_path)
    url = sheet.rows[0].value("Image")
    data = _edit(path, {("5106912", "Image"): f'=_xlfn.IMAGE("{url}")'})
    assert preview(data, "RFQ.xlsx", "job-1", sheet, 1).changes == []


def test_conflict_with_edit_made_in_app_after_download(tmp_path):
    from rfq.models import Cell, Sheet

    base, path = _download(tmp_path)
    current = Sheet.from_dict(base.to_dict())
    row = next(r for r in current.rows if r.product_id == "5106880")
    row.cells["ANSI"] = Cell("ANSI Z535.4", "manual_edit", "high", "colleague")  # changed in the app at v2
    data = _edit(path, {("5106880", "ANSI"): "ANSI Z535.1", ("5106880", "Temperature Range"): "-40 to 180 Deg F"})
    p = preview(data, "RFQ.xlsx", "job-1", current, 2, base=base)
    by = {(c.product_id, c.column): c for c in p.changes}
    assert by[("5106880", "ANSI")].status == "conflict" and "since your download" in by[("5106880", "ANSI")].message
    assert by[("5106880", "Temperature Range")].status == "ok"
    assert any("downloaded at version 1" in w for w in p.warnings)


def test_wrong_job_and_bad_files(tmp_path):
    sheet, path = _download(tmp_path, job_id="other-job")
    assert "different RFQ" in preview(path.read_bytes(), "RFQ.xlsx", "job-1", sheet, 1).errors[0]
    assert preview(b"not a zip", "RFQ.xlsx", "job-1", sheet, 1).errors
    assert "xlsx" in preview(b"a,b", "RFQ.csv", "job-1", sheet, 1).errors[0]
    renamed = _edit(path, {}, rename=("Weight", "Wt"))
    assert "missing these columns: Weight" in preview(renamed, "RFQ.xlsx", "other-job", sheet, 1).errors[0]


def test_unknown_duplicate_and_missing_rows(tmp_path):
    sheet, path = _download(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["RFQ"]
    ws.append(["9999999", "made up"] + [None] * 21)
    ws.delete_rows(3)  # 5106935 removed in Excel
    buf = io.BytesIO()
    wb.save(buf)
    p = preview(buf.getvalue(), "RFQ.xlsx", "job-1", sheet, 1)
    assert any("9999999 is not in this RFQ" in w for w in p.warnings)
    assert any("kept unchanged: 5106935" in w for w in p.warnings)
    ws.append(["5106912"] + [None] * 22)
    buf = io.BytesIO()
    wb.save(buf)
    assert any("more than once" in e for e in preview(buf.getvalue(), "RFQ.xlsx", "job-1", sheet, 1).errors)


def test_no_meta_compares_with_current(tmp_path):
    sheet, path = _download(tmp_path)
    data = _edit(path, {("5106880", "ANSI"): "ANSI Z535.1"}, drop_meta=True)
    p = preview(data, "RFQ.xlsx", "job-1", sheet, 1)
    assert p.ok and len(p.changes) == 1 and any("version stamp" in w for w in p.warnings)


def test_image_column_is_not_read_picture_is(tmp_path):
    sheet, path = _download(tmp_path)
    data = _edit(path, {("5106912", "Image"): "https://example.com/x.png",
                        ("5106935", "Picture"): "https://i.ferg-img.com/i/fergusonprod/new_primary"})
    p = preview(data, "RFQ.xlsx", "job-1", sheet, 1)
    assert [(c.product_id, c.column, c.status) for c in p.changes] == [("5106935", "Picture", "ok")]
