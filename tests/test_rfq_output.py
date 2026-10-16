"""Part A acceptance: email 1 equals the reference; emails 2-4 flag what is not found."""
from __future__ import annotations

import openpyxl
from helpers import REFERENCE, make_eml, run_sync


def _rows(path, sheet):
    return [list(r) for r in openpyxl.load_workbook(path)[sheet].iter_rows(values_only=True)]


def test_email1_matches_reference_cell_by_cell(offline_run):
    result = offline_run("email-1")
    assert _rows(result.xlsx_path, "RFQ") == _rows(REFERENCE, "RFQ")


def test_email1_sources_match_reference(offline_run):
    result = offline_run("email-1")
    ours = _rows(result.xlsx_path, "Sources")
    # Extra on our side: a 7th "link" column, Description rows (every value needs a source) and the
    # supplier rows ("Supplier: …"). Everything the reference has is identical and in the same order.
    rfq_rows = [r[:6] for r in ours if r[1] != "Description" and not str(r[1]).startswith("Supplier: ")]
    assert rfq_rows == _rows(REFERENCE, "Sources")
    assert ours[0][6] == "link"
    descriptions = [r for r in ours if r[1] == "Description"]
    assert len(descriptions) == 6 and all(r[3] and r[5] for r in descriptions)


def test_every_filled_value_has_a_source(offline_run):
    for name in ("email-1", "email-2", "email-3", "email-4"):
        result = offline_run(name)
        wb = openpyxl.load_workbook(result.xlsx_path)
        rfq = list(wb["RFQ"].iter_rows(values_only=True))
        header = rfq[0]
        sourced = {(str(r[0]), r[1]) for r in wb["Sources"].iter_rows(min_row=2, values_only=True)}
        for row in rfq[1:]:
            for column, value in zip(header[1:-1], row[1:-1]):
                if value not in (None, ""):
                    assert (str(row[0]), column) in sourced, (name, row[0], column)


def test_email2_offline_225123_not_found_kept_in_order(offline_run):
    result = offline_run("email-2")
    rows = _rows(result.xlsx_path, "RFQ")[1:]
    assert [r[0] for r in rows][-1] == "225123" and len(rows) == 8
    assert rows[-1][-1] == "Not found" and all(v is None for v in rows[-1][1:-1])
    assert [r[-1] for r in rows[:-1]] == ["Found"] * 7


def test_email3_competitors_not_found(offline_run):
    result = offline_run("email-3")
    status = {r[0]: r[-1] for r in _rows(result.xlsx_path, "RFQ")[1:]}
    assert [k for k, v in status.items() if v == "Found"] == ["5106878", "5106907", "4570880", "9677531"]
    assert [k for k, v in status.items() if v == "Not found"] == ["17783-79-3000", "31-024", "MTX-4409-XQ"]


def test_email4_duplicate_note_and_not_found(offline_run):
    result = offline_run("email-4")
    status = {r[0]: r[-1] for r in _rows(result.xlsx_path, "RFQ")[1:]}
    assert list(status).count("5106926") == 1
    assert status["9677533"] == "Not found" and status["17774-79-3000"] == "Not found"
    notes = _rows(result.xlsx_path, "Notes")
    assert any(n[0] == "5106926" and "2 times" in n[1] for n in notes)


def test_zero_ids_still_writes_rfq_and_summary(tmp_path):
    result = run_sync(tmp_path, make_eml("Hello, please quote some tape."))
    wb = openpyxl.load_workbook(result.xlsx_path)
    assert wb["RFQ"].max_row == 1
    text = " ".join(str(v) for r in wb["Summary"].iter_rows(values_only=True) for v in r if v)
    assert "No product IDs" in text
    assert "could not find any product IDs" in result.reply_path.read_text()


def test_reply_eml_has_attachment_and_lists_not_found(offline_run):
    import email
    import email.policy

    result = offline_run("email-3")
    msg = email.message_from_bytes(result.reply_path.read_bytes(), policy=email.policy.default)
    names = [p.get_filename() for p in msg.iter_attachments()]
    assert names == ["RFQ.xlsx"]
    body = msg.get_body(("plain",)).get_content()
    assert "4 of 7 products found" in body and "MTX-4409-XQ" in body


def test_formula_injection_is_written_as_text(tmp_path):
    raw = make_eml('=HYPERLINK("http://evil","x") MTX-4409-XQ\n+cmd 31-024')
    result = run_sync(tmp_path, raw)
    wb = openpyxl.load_workbook(result.xlsx_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":  # only our own picture formula may be a formula
                    assert str(cell.value).startswith('=IMAGE("https://'), (ws.title, cell.coordinate, cell.value)
    summary = [c for r in wb["Summary"].iter_rows() for c in r if str(c.value).startswith("=HYPERLINK")]
    assert summary and all(c.data_type == "s" for c in summary)  # kept as visible text


def test_cli_contract(tmp_path):
    import subprocess
    import sys

    from helpers import EMAILS, ROOT

    src = next(EMAILS.glob("email-1*.eml"))
    proc = subprocess.run([sys.executable, "rfq/build.py", str(src), "--out", str(tmp_path), "--offline",
                           "--no-llm", "--no-pdf", "--quiet"], cwd=ROOT, capture_output=True, text=True, timeout=120)
    assert proc.returncode == 0, proc.stderr
    assert (tmp_path / src.stem / "RFQ.xlsx").exists() and (tmp_path / src.stem / "reply.eml").exists()


def test_cli_missing_file(tmp_path):
    import subprocess
    import sys

    from helpers import ROOT

    proc = subprocess.run([sys.executable, "rfq/build.py", "nope.eml", "--out", str(tmp_path)], cwd=ROOT,
                          capture_output=True, text=True, timeout=60)
    assert proc.returncode == 2 and "not found" in proc.stderr
