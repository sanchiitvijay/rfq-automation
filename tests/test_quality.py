"""Part B acceptance: the Quality Agent flags, never overwrites; Summary + fills; graceful LLM failure."""
from __future__ import annotations

import asyncio
import hashlib
import json
from types import SimpleNamespace

import openpyxl
from helpers import run_sync

from rfq.config import RunOptions
from rfq.excel.reader import read
from rfq.excel.writer import render
from rfq.models import Cell, Row, Sheet
from rfq.quality.agent import run_quality
from rfq.quality.checks import run_code_checks


def _values_hash(path) -> str:
    ws = openpyxl.load_workbook(path)["RFQ"]
    return hashlib.sha256(json.dumps([list(r) for r in ws.iter_rows(values_only=True)], default=str).encode()).hexdigest()


def _row(**values) -> Row:
    row = Row(product_id="5106912", kind="ferguson", lookup_status="Found", lookup_source="catalog")
    for k, v in values.items():
        row.cells[k.replace("_", " ")] = Cell(v, "test", "high", "test")
    return row


def test_quality_never_changes_values(offline_run, tmp_path):
    result = offline_run("email-2")
    plain = render(result.sheet, tmp_path / "plain.xlsx")  # same sheet, no quality pass
    assert _values_hash(plain) == _values_hash(result.xlsx_path)


def test_quality_agent_on_finished_workbook_keeps_values(offline_run, tmp_path):
    result = offline_run("email-1")
    before = _values_hash(result.xlsx_path)
    sheet = read(result.xlsx_path)
    report = asyncio.run(run_quality(sheet, RunOptions(use_llm=False)))
    out = render(sheet, tmp_path / "checked.xlsx", report)
    assert _values_hash(out) == before


def test_summary_sheet_and_highlighting(offline_run):
    result = offline_run("email-2")
    wb = openpyxl.load_workbook(result.xlsx_path)
    assert wb.sheetnames[0] == "RFQ" and "Summary" in wb.sheetnames  # opens on the quotation itself
    text = [str(v) for r in wb["Summary"].iter_rows(values_only=True) for v in r if v is not None]
    for heading in ("Fill rate per column", "Fill rate per product", "Not found (left blank, nothing guessed)",
                    "Missing fields (no source had a value)"):
        assert heading in text
    assert "225123" in text
    fills = {c.fill.fgColor.rgb[-6:] for r in wb["RFQ"].iter_rows(min_row=2) for c in r if c.fill.fill_type}
    assert "D9D9D9" in fills  # missing (grey)
    assert "FBE2D5" in fills  # not-found row


def test_email2_missing_weights_flagged(offline_run):
    flags = offline_run("email-2").report.flags
    assert any(f.column == "Weight" and f.type == "missing" for f in flags)


def test_unit_conversion_flag():
    flags = run_code_checks(Sheet(rows=[_row(Length="333.33 yds.", Width="6")]))
    reasons = {f.column: f.reason for f in flags if f.type == "suspicious"}
    assert "333.33 yds. = 1000 ft" in reasons["Length"]
    assert "no unit given" in reasons["Width"]


def test_contradictions_by_code():
    row = _row(Description="3 in x 1000 ft Red Tape", Width="6 in", Length="1000 ft", Color_Finish="Blue",
               Lettering="Caution Buried Gas Line Below", Application="Water|Marking")
    flags = {(f.column, f.type) for f in run_code_checks(Sheet(rows=[row]))}
    assert ("Width", "suspicious") in flags  # description 3 in vs 6 in
    assert ("Color Finish", "suspicious") in flags  # blue ≠ gas (APWA yellow) and ≠ "red" in description
    assert ("Application", "suspicious") in flags  # water vs gas lettering


def test_email1_application_vs_lettering(offline_run):
    flags = offline_run("email-1").report.flags
    hit = [f for f in flags if f.product_id == "5106891" and f.column == "Application"]
    assert hit and hit[0].cell == "D6" and "sewer" in hit[0].reason


def test_plausibility_ranges():
    flags = run_code_checks(Sheet(rows=[_row(Width="300 in", Thickness="0.2 mil")]))
    assert {f.column for f in flags if "outside" in f.reason} == {"Width", "Thickness"}


def test_not_found_row_one_flag_not_twenty():
    row = Row(product_id="MTX-4409-XQ", kind="competitor", not_found_reason="competitor code")
    flags = run_code_checks(Sheet(rows=[row]))
    assert len(flags) == 1 and flags[0].type == "not_found" and flags[0].cell == "A2"


class _FakeCompletions:
    def __init__(self, content=None, error=None):
        self.content, self.error, self.calls = content, error, 0

    async def create(self, **kwargs):
        self.calls += 1
        if self.error:
            raise self.error
        assert "<rows>" in kwargs["messages"][1]["content"]
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=self.content))])


def _client(content=None, error=None):
    completions = _FakeCompletions(content, error)
    return SimpleNamespace(chat=SimpleNamespace(completions=completions)), completions


def test_llm_failure_degrades_to_code_checks(tmp_path):
    from helpers import eml

    client, _ = _client(error=TimeoutError("model timed out"))
    result = run_sync(tmp_path, eml("email-1"), RunOptions(offline=True, use_llm=True, use_pdfs=False),
                      llm_client=client)
    assert result.report.llm_status.startswith("skipped: TimeoutError")
    assert any(f.check == "code" for f in result.report.flags)
    summary = [str(v) for r in openpyxl.load_workbook(result.xlsx_path)["Summary"].iter_rows(values_only=True)
               for v in r if v]
    assert any("skipped: TimeoutError" in s for s in summary)


def test_llm_flags_validated_and_cached():
    answer = json.dumps({"flags": [
        {"product_id": "5106912", "column": "Lettering", "reason": "Lettering says gas, colour is blue"},
        {"product_id": "5106912", "column": "OSHA", "reason": "blank column, must be ignored"},
        {"product_id": "9999999", "column": "Width", "reason": "unknown product, must be ignored"},
        {"product_id": "5106912", "column": "Price", "reason": "not an RFQ column"},
    ]})
    client, calls = _client(answer)
    sheet = Sheet(rows=[_row(Lettering="Caution Buried Gas Line Below", Color_Finish="Blue")])
    from rfq.lookup.cache import MemoryCache

    cache = MemoryCache()
    report = asyncio.run(run_quality(sheet, RunOptions(use_llm=True), cache, client))
    llm = [f for f in report.flags if f.check == "llm" or "colour is blue" in f.reason]
    assert len(llm) == 1 and llm[0].column == "Lettering" and llm[0].cell == "G2"
    asyncio.run(run_quality(sheet, RunOptions(use_llm=True), cache, client))
    assert calls.calls == 1  # second run served from cache


def test_llm_bad_json_is_rejected():
    client, _ = _client("not json at all")
    report = asyncio.run(run_quality(Sheet(rows=[_row(Width="6 in")]), RunOptions(use_llm=True), llm_client=client))
    assert "did not match the schema" in report.llm_status


def test_llm_missing_key(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    report = asyncio.run(run_quality(Sheet(rows=[_row(Width="6 in")]), RunOptions(use_llm=True)))
    assert "OPENAI_API_KEY is not set" in report.llm_status


def test_fill_rates(offline_run):
    report = offline_run("email-3").report
    assert report.fill_by_product["MTX-4409-XQ"] == 0.0
    assert 0 < report.fill_overall < report.fill_found <= 100
    assert set(report.fill_by_column) and all(0 <= v <= 100 for v in report.fill_by_column.values())
