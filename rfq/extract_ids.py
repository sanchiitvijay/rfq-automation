"""Finds every requested product in the email body and spreadsheet attachments.

Deterministic on purpose (no LLM): every token returned appears verbatim in the source text,
so nothing can be invented here.
- Ferguson IDs: 6-7 digit numbers, optionally written "#5106926".
- Competitor / unknown codes: part numbers with a dash such as "17783-79-3000" or "MTX-4409-XQ",
  or upper-case letter+digit codes such as "SD6105B52".
Numbers that only look like IDs (dates, phone numbers, PO/order numbers, e-mail addresses, URLs)
are masked before scanning.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field

from rfq.config import MAX_IDS
from rfq.models import RequestItem

FERGUSON_ID = re.compile(r"(?<![\w\-/.])#?(\d{6,7})(?![\w\-/]|\.\d)")
DASH_CODE = re.compile(r"(?<![\w\-/#.])([A-Za-z0-9]+(?:-[A-Za-z0-9]+)+)(?![\w\-/])")
ALNUM_CODE = re.compile(r"(?<![\w\-])([A-Z0-9]{6,20})(?![\w\-])")

_MASKS = [
    re.compile(r"\S+@\S+"),  # e-mail addresses
    re.compile(r"https?://\S+|www\.\S+"),  # links
    re.compile(r"\b\d{4}-\d{2}-\d{2}\b|\b\d{1,2}[/.]\d{1,2}[/.]\d{2,4}\b"),  # dates
    re.compile(r"(?:\+\d{1,3}[\s.-]?)?(?:\(\d{3}\)|\d{3})[\s.-]\d{3}[\s.-]\d{4}\b"),  # phone numbers
    re.compile(r"\b1-800-\d{3}-\d{4}\b"),
    re.compile(
        r"\b(?:P\.?O\.?|order|invoice|ticket|ref(?:erence)?|zip|phone|tel|fax|mobile|cell|acct|account|"
        r"case|quote|job)\b\s*(?:#|no\.?|number|num|:)?\s*:?\s*#?\d[\d-]*",
        re.IGNORECASE,
    ),
]
_QTY = re.compile(
    r"(?:\bx\s*|\bqty\.?\s*:?\s*|\bquantity\s*:?\s*)(\d{1,6})\b|\b(\d{1,6})\s*(?:rolls?|pcs?|pieces|units?|ea|each|boxes|cases)\b",
    re.IGNORECASE,
)
_ID_HEADER = re.compile(r"item|product|part|sku|catalog|article|^\s*id\s*$|code", re.IGNORECASE)
_QTY_HEADER = re.compile(r"qty|quantity|count|rolls", re.IGNORECASE)

KNOWN_BRANDS = [
    "Mutual Industries", "Mutual", "Empire", "Presco", "Pro-Pak Industries", "Pro-Pak", "Pro-Line",
    "King Sales", "Paddock", "Northtown", "Rhino", "Reef Industries", "Christy's", "Seton", "Brady",
    "3M", "Harris", "Terra Tape", "Calpico", "Brooks",
]


@dataclass
class ExtractResult:
    items: list[RequestItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _mask(line: str) -> str:
    for pattern in _MASKS:
        line = pattern.sub(lambda m: " " * len(m.group(0)), line)
    return line


def _is_competitor_code(token: str) -> bool:
    digits = sum(ch.isdigit() for ch in token)
    if digits < 3 or ("-" in token and digits < 4):  # "3-inch", "Q4-26", "Fwd-2"
        return False
    if re.fullmatch(r"\d{6,7}", token):
        return False
    if "-" not in token:
        letters = sum(ch.isalpha() for ch in token)
        return token.isupper() and letters >= 2 and digits >= 3
    return True


def _brand_in(line: str) -> str | None:
    for brand in KNOWN_BRANDS:
        if re.search(rf"(?<!\w){re.escape(brand)}(?!\w)", line, re.IGNORECASE):
            return brand
    return None


def _qty_in(line: str) -> str | None:
    m = _QTY.search(line)
    return (m.group(1) or m.group(2)) if m else None


def _scan_line(line: str) -> list[tuple[int, str, str, str]]:
    """(offset, raw_token, product_id, kind) for every product token on one line."""
    masked = _mask(line)
    found: list[tuple[int, str, str, str]] = []
    taken: list[range] = []
    for m in DASH_CODE.finditer(masked):
        token = m.group(1)
        if _is_competitor_code(token):
            found.append((m.start(1), token, token, "competitor"))
            taken.append(range(m.start(1), m.end(1)))
    for m in FERGUSON_ID.finditer(masked):
        if any(m.start() in r for r in taken):
            continue
        found.append((m.start(), m.group(0), m.group(1), "ferguson"))
        taken.append(range(m.start(), m.end()))
    for m in ALNUM_CODE.finditer(masked):
        token = m.group(1)
        if any(m.start(1) in r for r in taken) or token.isdigit() or not _is_competitor_code(token):
            continue
        found.append((m.start(1), token, token, "competitor"))
    return sorted(found)


class _Collector:
    def __init__(self) -> None:
        self.items: list[RequestItem] = []
        self.by_key: dict[str, RequestItem] = {}

    def add(self, raw: str, pid: str, kind: str, context: str, origin: str, qty: str | None) -> None:
        key = pid.upper()
        if key in self.by_key:
            self.by_key[key].occurrences += 1
            return
        item = RequestItem(
            position=len(self.items) + 1, raw_token=raw, product_id=pid, kind=kind,  # type: ignore[arg-type]
            context=context.strip()[:300], origin=origin, qty=qty,
            brand=_brand_in(context) if kind == "competitor" else None,
        )
        self.items.append(item)
        self.by_key[key] = item


def _scan_text(text: str, collector: _Collector, origin: str) -> None:
    for line in text.splitlines():
        hits = _scan_line(line)
        qty = _qty_in(_mask(line)) if len(hits) == 1 else None
        for _, raw, pid, kind in hits:
            collector.add(raw, pid, kind, line, origin, qty)


def _classify_cell(value: str) -> tuple[str, str] | None:
    value = value.strip()
    if re.fullmatch(r"\d{6,7}\.0", value):  # numbers stored as floats in Excel
        value = value[:-2]
    m = re.fullmatch(r"#?(\d{6,7})", value)
    if m:
        return m.group(1), "ferguson"
    if re.fullmatch(r"[A-Za-z0-9]+(?:-[A-Za-z0-9]+)*", value) and _is_competitor_code(value.upper() if "-" not in value else value):
        return value, "competitor"
    return None


def _table_rows(name: str, data: bytes, warnings: list[str]) -> list[list[list[str]]]:
    """Every sheet of the attachment as rows of strings."""
    lower = name.lower()
    if lower.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [[[c for c in row] for row in csv.reader(io.StringIO(text))]]
    if lower.endswith(".xls"):
        warnings.append(f"'{name}' is a legacy .xls file and was not read; please resend it as .xlsx.")
        return []
    from openpyxl import load_workbook  # local import keeps CLI start-up fast

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # values only, never formulas/macros
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as e:
        warnings.append(f"'{name}' could not be opened (corrupt or password protected): {type(e).__name__}.")
        return []
    sheets = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        sheets.append(rows)
    wb.close()
    return sheets


def _scan_attachment(name: str, data: bytes, collector: _Collector, warnings: list[str]) -> None:
    origin = f"attachment:{name}"
    sheets = _table_rows(name, data, warnings)
    before = len(collector.items)
    for rows in sheets:
        if not any(any(c.strip() for c in r) for r in rows):
            continue
        header_idx, id_col, qty_col = None, None, None
        for i, row in enumerate(rows[:10]):
            for j, cell in enumerate(row):
                if id_col is None and _ID_HEADER.search(cell) and not _QTY_HEADER.search(cell):
                    header_idx, id_col = i, j
                elif _QTY_HEADER.search(cell):
                    qty_col = j
            if id_col is not None:
                break
        if id_col is None:  # no recognisable header: scan every cell like body text
            _scan_text("\n".join(" ".join(r) for r in rows), collector, origin)
            continue
        for row in rows[header_idx + 1:]:  # type: ignore[operator]
            if id_col >= len(row) or not row[id_col].strip():
                continue
            hit = _classify_cell(row[id_col])
            qty = row[qty_col].strip() if qty_col is not None and qty_col < len(row) else None
            if hit is None:
                warnings.append(f"'{name}': value '{row[id_col][:40]}' in the item column is not a product ID; skipped.")
                continue
            collector.add(row[id_col].strip(), hit[0], hit[1], " | ".join(c for c in row if c), origin, qty or None)
    if len(collector.items) == before and sheets:
        warnings.append(f"No product IDs found in attachment '{name}'.")


def extract(body: str, attachments: list[tuple[str, bytes]] | None = None) -> ExtractResult:
    result = ExtractResult()
    collector = _Collector()
    _scan_text(body, collector, "body")
    for name, data in attachments or []:
        _scan_attachment(name, data, collector, result.warnings)
    items = collector.items
    if len(items) > MAX_IDS:
        result.warnings.append(f"{len(items)} products requested; only the first {MAX_IDS} were processed.")
        items = items[:MAX_IDS]
    if not items:
        result.warnings.append("No product IDs were found in the email.")
    result.items = items
    return result
