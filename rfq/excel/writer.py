"""Renders a Sheet (+ optional QualityReport) into the RFQ template.

The workbook is always generated from state, never edited in place, so the web editor and the file
cannot drift apart. External text is always written as a string cell: only our own =IMAGE() formula
is ever a formula (no formula injection from emails, web pages or PDFs).
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell as XlCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from rfq.config import (
    NOTES_COLUMNS, RFQ_COLUMNS, SOURCES_COLUMNS, STATUS_NOT_FOUND, SUPPLIER_COLUMNS, SUPPLIER_PREFIX, TEMPLATE_PATH,
)
from rfq.excel.images import primary_picture
from rfq.models import Cell, QualityReport, Sheet

from rfq.excel.palette import CONFIDENCE_COLOUR, STATUS_COLOUR, fill_of, text_of  # noqa: E402


def _fill(name: str) -> PatternFill:
    return PatternFill("solid", fgColor=fill_of(name))


GREY = _fill("missing")
ORANGE = _fill("suspicious")
MANUAL = _fill("manual")
NOT_FOUND = _fill("not_found")
HEADER_FILL = _fill("header")
HEADER_FONT = Font(bold=True, color=text_of("header"))
SAFE_URL = re.compile(r"^https?://[^\s\"<>|]+$")
LINK_FONT = Font(color="1379C8", underline="single")
LINK_COLUMNS = {"Product page", "Spec sheet"}


def column_letter(column: str) -> str:
    return get_column_letter(RFQ_COLUMNS.index(column) + 1)


def _put(cell: XlCell, value: str) -> None:
    """Writes external text as a plain string, never as a formula."""
    cell.value = value
    if isinstance(value, str) and value:
        cell.data_type = "s"


def _image_formula(url: str) -> str | None:
    return f'=IMAGE("{url}")' if SAFE_URL.match(url) else None


def _header(ws: Worksheet, columns: list[str]) -> None:
    for j, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL


def _fill_rfq(ws: Worksheet, sheet: Sheet) -> None:
    for i, row in enumerate(sheet.rows, start=2):
        for j, column in enumerate(RFQ_COLUMNS, start=1):
            target = ws.cell(row=i, column=j)
            if column == "product_id":
                _put(target, row.product_id)
            elif column == "Lookup Status":
                _put(target, row.lookup_status)
            elif column == "Image":
                url = row.value("Image")
                formula = _image_formula(primary_picture(url)) if "|" in url else (_image_formula(url) if url else None)
                if formula:
                    target.value = formula  # our own formula, URL validated above
                elif url:
                    _put(target, url)
            elif column == "Picture":
                url = row.value("Picture")
                if SAFE_URL.match(url):
                    _link(target, url)  # clickable; the picture itself is in the Image column
                elif url:
                    _put(target, url)  # e.g. several URLs separated by |: kept as written
                    first = primary_picture(url)
                    if first:
                        target.hyperlink = first
            else:
                value = row.value(column)
                if value:
                    _put(target, value)
            target.alignment = Alignment(vertical="top", wrap_text=column == "Product details")


def _link(target: XlCell, url: str) -> None:
    """A clickable link, only for plain http(s) addresses."""
    if url and SAFE_URL.match(url):
        _put(target, url)
        target.hyperlink = url
        target.font = LINK_FONT


IMAGE_COLUMN_WIDTH = 13.0  # ~96 px, a little narrower than the 100 px picture so it covers the cell
HIDDEN_FALLBACK = Font(color="FFFFFF")
IMAGE_ROW_HEIGHT = 77  # points (~102 px)


def _embed_pictures(ws: Worksheet, sheet: Sheet) -> int:
    """Puts each product's thumbnail into its Image cell, so the picture shows in every spreadsheet app.
    The cell keeps the =IMAGE() formula underneath (Excel 365 draws it too; the picture covers it)."""
    from openpyxl.drawing.image import Image as XlImage

    from rfq.excel import images

    col = column_letter("Image")
    embedded = 0
    for i, row in enumerate(sheet.rows, start=2):
        url = images.primary_picture(row.value("Picture") or row.value("Image"))
        path = images.cached(url) if url else None
        if path is None:
            continue
        pic = XlImage(str(path))
        pic.width = pic.height = images.THUMB_PX
        ws.add_image(pic, f"{col}{i}")
        ws.row_dimensions[i].height = IMAGE_ROW_HEIGHT
        # apps without IMAGE() show "0" or #NAME? for the formula: keep that text hidden under the picture
        cell = ws[f"{col}{i}"]
        cell.font = HIDDEN_FALLBACK
        cell.alignment = Alignment(horizontal="left", vertical="top")
        embedded += 1
    if embedded:
        ws.column_dimensions[col].width = IMAGE_COLUMN_WIDTH
    return embedded


def _fill_sources(ws: Worksheet, sheet: Sheet) -> None:
    """One row per filled value: RFQ columns first, then the product's supplier values."""
    _header(ws, SOURCES_COLUMNS)
    r = 2

    def write(pid: str, column: str, cell: Cell) -> None:
        nonlocal r
        for j, value in enumerate([pid, column, cell.value, cell.source_type, cell.confidence, cell.source],
                                  start=1):
            _put(ws.cell(row=r, column=j), value)
        _link(ws.cell(row=r, column=7), cell.url)
        conf = CONFIDENCE_COLOUR.get((cell.confidence or "").lower())
        if conf:
            ws.cell(row=r, column=5).fill = _fill(conf)
            ws.cell(row=r, column=5).font = Font(color=text_of(conf))
        if cell.source_type == "manual_edit":
            ws.cell(row=r, column=4).fill = MANUAL
        r += 1

    for row in sheet.rows:
        for column in RFQ_COLUMNS[1:-1]:
            cell = row.cells.get(column)
            if cell and cell.value:
                write(row.product_id, column, cell)
        for column in SUPPLIER_COLUMNS:
            cell = row.supplier.get(column)
            if cell and cell.value:
                write(row.product_id, SUPPLIER_PREFIX + column, cell)
    for col, width in {"A": 14, "B": 26, "C": 40, "D": 20, "E": 11, "F": 70, "G": 60}.items():
        ws.column_dimensions[col].width = width


def _fill_supplier(ws: Worksheet, sheet: Sheet) -> None:
    columns = ["product_id", *SUPPLIER_COLUMNS, "Lookup Status", "Found via"]
    _header(ws, columns)
    for i, row in enumerate(sheet.rows, start=2):
        for j, column in enumerate(columns, start=1):
            target = ws.cell(row=i, column=j)
            if column == "product_id":
                _put(target, row.product_id)
            elif column == "Lookup Status":
                _put(target, row.lookup_status)
                _status_style(target, row.lookup_status)
            elif column == "Found via":
                _put(target, {"live": "ferguson.com page", "cache": "ferguson.com page (cached)",
                              "serpapi+catalog": "SerpApi + backup catalog", "serpapi": "SerpApi (Google)",
                              "catalog": "backup catalog", "manual": "entered by hand", "none": ""}
                     .get(row.lookup_source, row.lookup_source))
            else:
                cell = row.supplier.get(column)
                if not cell or not cell.value:
                    continue
                if column in LINK_COLUMNS:
                    _link(target, cell.value)
                else:
                    _put(target, cell.value)
        if row.lookup_status == STATUS_NOT_FOUND:
            for j in range(1, len(columns) + 1):
                if columns[j - 1] != "Lookup Status":
                    ws.cell(row=i, column=j).fill = NOT_FOUND
    note_row = len(sheet.rows) + 3
    _put(ws.cell(row=note_row, column=1), "Where each value comes from (with a link) is on the Sources sheet, "
                                          f"under '{SUPPLIER_PREFIX}<column>'. Catalog prices date from 2026-09-24.")
    ws.freeze_panes = "B2"
    for col, width in {"A": 14, "B": 12, "C": 20, "D": 20, "E": 18, "F": 15, "G": 12, "H": 14, "I": 60,
                       "J": 60, "K": 13, "L": 24}.items():
        ws.column_dimensions[col].width = width


def _fill_notes(ws: Worksheet, sheet: Sheet) -> None:
    _header(ws, NOTES_COLUMNS)
    r = 2
    for warning in sheet.warnings:
        _put(ws.cell(row=r, column=1), "(all)")
        _put(ws.cell(row=r, column=2), warning)
        r += 1
    for note in sheet.notes:
        _put(ws.cell(row=r, column=1), note.product_id)
        _put(ws.cell(row=r, column=2), note.note)
        r += 1
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110


def _status_style(cell: XlCell, status: str) -> None:
    name = STATUS_COLOUR.get(status)
    if name:
        cell.fill = _fill(name)
        cell.font = Font(bold=True, color=text_of(name))


def _comment(cell: XlCell, text: str) -> None:
    from openpyxl.comments import Comment

    note = Comment(text[:1000], "RFQ Agent")
    note.width, note.height = 320, 120
    cell.comment = note


def _highlight(ws: Worksheet, sheet: Sheet, report: QualityReport) -> None:
    """Same colours as the web app. Order matters: edited-by-hand first, flags on top (a flag is more urgent)."""
    row_of = {r.product_id: i for i, r in enumerate(sheet.rows, start=2)}
    for i, row in enumerate(sheet.rows, start=2):
        for column, cell in row.cells.items():
            if cell.source_type == "manual_edit" and column in RFQ_COLUMNS:
                ws[f"{column_letter(column)}{i}"].fill = MANUAL
        _status_style(ws[f"{column_letter('Lookup Status')}{i}"], row.lookup_status)
    reasons: dict[str, list[str]] = {}
    for flag in report.flags:
        if flag.type == "not_found":
            i = row_of.get(flag.product_id)
            if i:
                for j in range(1, len(RFQ_COLUMNS)):  # whole row except the coloured status cell
                    ws.cell(row=i, column=j).fill = NOT_FOUND
                _comment(ws.cell(row=i, column=1), f"Not found: {flag.reason}")
            continue
        ws[flag.cell].fill = GREY if flag.type == "missing" else ORANGE
        if flag.type == "suspicious":
            label = "AI check" if flag.check == "llm" else "Rule"
            reasons.setdefault(flag.cell, []).append(f"{label}: {flag.reason}")
    for ref, lines in reasons.items():  # hover the cell in Excel to read why it is orange
        _comment(ws[ref], "\n".join(lines))


META_SHEET = "_rfq_meta"


def _write_meta(wb, meta: dict[str, str]) -> None:
    """Hidden sheet so an edited copy that is uploaded again can be matched to its job and version."""
    ws = wb.create_sheet(META_SHEET)
    for i, (key, value) in enumerate(meta.items(), start=1):
        _put(ws.cell(row=i, column=1), key)
        _put(ws.cell(row=i, column=2), str(value))
    ws.sheet_state = "hidden"


def render(sheet: Sheet, path: Path, report: QualityReport | None = None, email_info: dict | None = None,
           meta: dict[str, str] | None = None, embed_images: bool = True) -> Path:
    from rfq.quality.summary import write_summary  # summary depends on writer helpers

    wb = load_workbook(TEMPLATE_PATH)
    _fill_rfq(wb["RFQ"], sheet)
    _fill_sources(wb["Sources"], sheet)
    _fill_supplier(wb.create_sheet("Supplier", index=wb.sheetnames.index("RFQ") + 1), sheet)
    notes = wb["Notes"] if "Notes" in wb.sheetnames else wb.create_sheet("Notes")
    _fill_notes(notes, sheet)
    if report is not None:
        _highlight(wb["RFQ"], sheet, report)
        write_summary(wb, sheet, report, email_info or {})
    if "How to fill" in wb.sheetnames:  # keep the instructions sheet last
        wb.move_sheet("How to fill", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("How to fill"))
    if meta:
        _write_meta(wb, meta)
    ws = wb["RFQ"]
    # open on the RFQ (the quotation itself), not on the Summary
    for other in wb.worksheets:
        other.sheet_view.tabSelected = other.title == "RFQ"
    wb.active = wb.sheetnames.index("RFQ")
    ws.freeze_panes = "C2"
    for name in ("Sources", "Notes"):
        if name in wb.sheetnames:
            wb[name].freeze_panes = "A2"
    for column, width in {"A": 14, "B": 42, "R": 60, "U": 14, "V": 30}.items():
        ws.column_dimensions[column].width = width
    if embed_images:
        _embed_pictures(ws, sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)  # atomic: a reader never sees a half-written file
    return path


def render_summary_bytes(sheet: Sheet, report: QualityReport, email_info: dict | None = None) -> bytes:
    """A small workbook with only the Summary and the Notes (the 'download summary' button)."""
    import io

    from openpyxl import Workbook

    from rfq.quality.summary import write_summary

    wb = Workbook()
    write_summary(wb, sheet, report, email_info or {})
    del wb[wb.sheetnames[0]]  # the empty default sheet
    _fill_notes(wb.create_sheet("Notes"), sheet)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
