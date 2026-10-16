"""Reads a finished RFQ workbook back into a Sheet, so the Quality Agent can run on any RFQ.xlsx."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from rfq.config import RFQ_COLUMNS, STATUS_NOT_FOUND, SUPPLIER_PREFIX
from rfq.models import Cell, Note, Row, Sheet

_IMAGE = re.compile(r'^=IMAGE\("([^"]+)"', re.IGNORECASE)


def read(path: Path) -> Sheet:
    wb = load_workbook(path)
    ws = wb["RFQ"]
    header = [str(c.value or "").strip() for c in ws[1]]
    missing = [c for c in RFQ_COLUMNS if c not in header]
    if missing:
        raise ValueError(f"{path.name} is not an RFQ workbook (missing columns: {', '.join(missing)})")
    index = {name: i for i, name in enumerate(header)}

    sources: dict[tuple[str, str], Cell] = {}
    if "Sources" in wb.sheetnames:
        for values in wb["Sources"].iter_rows(min_row=2, max_col=7, values_only=True):
            pid, column, value, stype, conf, src = values[:6]
            link = values[6] if len(values) > 6 else None
            if pid is not None and column:
                sources[(str(pid), str(column))] = Cell(str(value or ""), str(stype or ""), str(conf or ""),
                                                        str(src or ""), str(link or ""))

    sheet = Sheet()
    for values in ws.iter_rows(min_row=2, values_only=True):
        pid = values[index["product_id"]]
        if pid is None or str(pid).strip() == "":
            continue
        pid = str(pid).strip()
        status = str(values[index["Lookup Status"]] or STATUS_NOT_FOUND)
        is_code = not re.fullmatch(r"\d{6,7}", pid)
        row = Row(product_id=pid, kind="competitor" if is_code else "ferguson", lookup_status=status)
        for column in RFQ_COLUMNS[1:-1]:
            raw = values[index[column]]
            if raw is None or str(raw).strip() == "":
                continue
            value = str(raw).strip()
            m = _IMAGE.match(value) if column == "Image" else None
            if m:
                value = m.group(1)
            src = sources.get((pid, column))
            row.cells[column] = Cell(value, src.source_type if src else "", src.confidence if src else "",
                                     src.source if src else "", src.url if src else "")
        for (spid, scol), cell in sources.items():  # supplier values live on the Sources sheet too
            if spid == pid and scol.startswith(SUPPLIER_PREFIX):
                row.supplier[scol[len(SUPPLIER_PREFIX):]] = cell
        sheet.rows.append(row)
    if "Notes" in wb.sheetnames:
        for pid, note in wb["Notes"].iter_rows(min_row=2, max_col=2, values_only=True):
            if note:
                sheet.notes.append(Note(str(pid or ""), str(note)))
    return sheet
