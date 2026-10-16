"""The Summary sheet: fill rates, what was not found and why, missing fields, and every flag."""
from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from rfq.config import STATUS_NOT_FOUND
from rfq.excel.palette import COLOURS, LEGEND as PALETTE_LEGEND
from rfq.models import QualityReport, Sheet

TITLE = Font(bold=True, size=14, color="0A2433")
SECTION = Font(bold=True, size=12, color="00446A")
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="F0F3F5")
LEGEND = {k: COLOURS[k][0] for k in ("missing", "suspicious", "not_found")}  # flag type → fill


class _Writer:
    def __init__(self, ws: Worksheet) -> None:
        self.ws, self.r = ws, 1

    def line(self, *values: object, font: Font | None = None, fill: PatternFill | None = None) -> None:
        for j, value in enumerate(values, start=1):
            cell = self.ws.cell(row=self.r, column=j, value=value)
            if isinstance(value, str) and value:
                cell.data_type = "s"  # external text, never a formula
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
        self.r += 1

    def head(self, *values: str) -> None:
        self.line(*values, font=BOLD, fill=HEAD_FILL)

    def section(self, title: str) -> None:
        self.r += 1
        self.line(title, font=SECTION)


def write_summary(wb: Workbook, sheet: Sheet, report: QualityReport, email_info: dict) -> None:
    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    after = wb.sheetnames.index("Supplier") + 1 if "Supplier" in wb.sheetnames else 1
    ws = wb.create_sheet("Summary", index=after)  # the RFQ itself stays the first sheet
    w = _Writer(ws)
    found = [r for r in sheet.rows if r.lookup_status != STATUS_NOT_FOUND]
    not_found = [r for r in sheet.rows if r.lookup_status == STATUS_NOT_FOUND]
    flags_by_type = {t: sum(1 for f in report.flags if f.type == t) for t in ("not_found", "suspicious", "missing")}

    w.line("RFQ quality summary", font=TITLE)
    if email_info.get("subject"):
        w.line("Request", email_info["subject"])
    if email_info.get("sender"):
        w.line("From", email_info["sender"])
    w.line("Generated", datetime.now().strftime("%Y-%m-%d %H:%M"))
    w.line("Lookup mode", "live ferguson.com + backup catalog" if sheet.mode == "live" else "offline (backup catalog + spec sheets)")
    w.line("AI contradiction check", report.llm_status)
    via: dict[str, int] = {}
    labels = {"live": "ferguson.com page", "cache": "ferguson.com page (cached)", "serpapi+catalog":
              "SerpApi + backup catalog", "serpapi": "SerpApi only", "catalog": "backup catalog",
              "manual": "entered by hand"}
    for r in found:
        via[labels.get(r.lookup_source, r.lookup_source)] = via.get(labels.get(r.lookup_source, r.lookup_source), 0) + 1
    if via:
        w.line("Found via", ", ".join(f"{k}: {v}" for k, v in via.items()))
    w.line("Supplier details", "Brand, part numbers, price, stock and product links are on the Supplier sheet; "
                               "every value's source and link is on the Sources sheet.")
    for warning in sheet.warnings:
        w.line("Warning", warning)

    w.section("Totals")
    w.line("Products requested", len(sheet.rows))
    w.line("Found", len(found))
    w.line("Not found", len(not_found))
    w.line("Fill rate, all requested products", f"{report.fill_overall}%")
    w.line("Fill rate, found products only", f"{report.fill_found}%")
    w.line("Flags: not found / suspicious / missing",
           f"{flags_by_type['not_found']} / {flags_by_type['suspicious']} / {flags_by_type['missing']}")
    if not sheet.rows:
        w.line("Result", "No product IDs were found in the email, so the RFQ sheet is empty.")

    w.section("Not found (left blank, nothing guessed)")
    if not_found:
        w.head("product_id", "Written in the email as", "Line in the email", "Why", "Sources tried")
        for r in not_found:
            w.line(r.product_id, r.raw_token, r.context, r.not_found_reason, "; ".join(r.tried))
    else:
        w.line("Every requested product was found.")

    w.section("Missing fields (no source had a value)")
    w.head("product_id", "Missing count", "Columns")
    for r in found:
        missing = r.missing_columns()
        if missing:
            w.line(r.product_id, len(missing), ", ".join(missing))

    w.section("Fill rate per product")
    w.head("product_id", "Lookup Status", "Source", "Fill rate")
    for r in sheet.rows:
        w.line(r.product_id, r.lookup_status, r.lookup_source, f"{report.fill_by_product.get(r.product_id, 0)}%")

    w.section("Fill rate per column")
    w.head("Column", "Fill rate")
    for column, rate in report.fill_by_column.items():
        w.line(column, f"{rate}%")

    w.section("Flags (cell values are never changed; review and fix by hand)")
    w.head("Cell (RFQ sheet)", "product_id", "Column", "Type", "Check", "Reason")
    for f in report.flags:
        w.line(f.cell, f.product_id, f.column, f.type, f.check, f.reason,
               fill=None)
        ws.cell(row=w.r - 1, column=4).fill = PatternFill("solid", fgColor=LEGEND[f.type])

    w.section("Colours used in this workbook")
    for name in PALETTE_LEGEND:
        fill, text, meaning = COLOURS[name]
        w.line(meaning, font=Font(color=text) if text else None, fill=PatternFill("solid", fgColor=fill))

    for col, width in {"A": 34, "B": 22, "C": 40, "D": 60, "E": 50, "F": 80}.items():
        ws.column_dimensions[col].width = width
